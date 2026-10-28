r"""
Batch Converter: Converts DOCX, PDF, XLSX, and MD into clean Markdown (.md)
Source: C:\Users\Saturn\Desktop\New folder\staged-sources-for-download
Destination: C:\Users\Saturn\Desktop\New folder\sources-for-notebook
Pure Python implementation - zero blocked DLLs.
"""

import os
import sys
import time
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
import openpyxl
from pypdf import PdfReader

SOURCE_DIR = Path(r"C:\Users\Saturn\Desktop\New folder\staged-sources-for-download")
DEST_DIR = Path(r"C:\Users\Saturn\Desktop\New folder\sources-for-notebook")
DEST_DIR.mkdir(parents=True, exist_ok=True)

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W_NS}

def get_target_filename(p: Path) -> str:
    name = p.name
    # Strip compound extensions in priority order
    for ext in ['.md.docx', '.pdf.docx', '.pdf.pdf', '.docx', '.pdf', '.xlsx', '.gdoc']:
        if name.lower().endswith(ext):
            base = name[:-len(ext)]
            return base + '.md'
    if name.lower().endswith('.md'):
        return name
    return p.stem + '.md'

def convert_docx(p: Path) -> str:
    if not zipfile.is_zipfile(str(p)):
        return p.read_text(encoding='utf-8', errors='replace')
    with zipfile.ZipFile(str(p)) as z:
        xml_content = z.read("word/document.xml")
    root = ET.fromstring(xml_content)
    body = root.find("w:body", NS)
    if body is None:
        return f"# {p.stem}\n\n*Empty document*"
    
    lines = [f"# {p.stem}\n"]
    for child in body:
        tag = child.tag
        if tag == f"{{{W_NS}}}p":
            # Paragraph
            pPr = child.find("w:pPr", NS)
            style_val = ""
            if pPr is not None:
                pStyle = pPr.find("w:pStyle", NS)
                if pStyle is not None:
                    style_val = pStyle.attrib.get(f"{{{W_NS}}}val", "").lower()
            
            texts = []
            for t in child.iter(f"{{{W_NS}}}t"):
                if t.text:
                    texts.append(t.text)
            para_text = "".join(texts).strip()
            if not para_text:
                continue
            
            if "heading1" in style_val or "title" in style_val:
                lines.append(f"\n# {para_text}\n")
            elif "heading2" in style_val or "subtitle" in style_val:
                lines.append(f"\n## {para_text}\n")
            elif "heading3" in style_val:
                lines.append(f"\n### {para_text}\n")
            elif "heading4" in style_val:
                lines.append(f"\n#### {para_text}\n")
            elif "list" in style_val or "bullet" in style_val:
                lines.append(f"* {para_text}")
            else:
                lines.append(para_text)
        elif tag == f"{{{W_NS}}}tbl":
            # Table
            rows = []
            for tr in child.findall("w:tr", NS):
                row_cells = []
                for tc in tr.findall("w:tc", NS):
                    cell_texts = []
                    for t in tc.iter(f"{{{W_NS}}}t"):
                        if t.text:
                            cell_texts.append(t.text)
                    cleaned = " ".join("".join(cell_texts).split()).replace("|", "\\|")
                    row_cells.append(cleaned)
                if any(row_cells):
                    rows.append(row_cells)
            if rows:
                num_cols = max(len(r) for r in rows)
                if num_cols > 0:
                    header = (rows[0] + [""] * num_cols)[:num_cols]
                    lines.append("\n| " + " | ".join(header) + " |")
                    lines.append("| " + " | ".join(["---"] * num_cols) + " |")
                    for r in rows[1:]:
                        cells = (r + [""] * num_cols)[:num_cols]
                        lines.append("| " + " | ".join(cells) + " |")
                    lines.append("\n")
    return "\n\n".join(lines)

def convert_xlsx(p: Path) -> str:
    wb = openpyxl.load_workbook(str(p), data_only=True)
    lines = [f"# Spreadsheet: {p.stem}\n"]
    for sheetname in wb.sheetnames:
        ws = wb[sheetname]
        lines.append(f"\n## Sheet: {sheetname}\n")
        rows = list(ws.iter_rows(values_only=True))
        rows = [r for r in rows if any(c is not None and str(c).strip() != '' for c in r)]
        if not rows:
            lines.append("*Empty sheet*\n")
            continue
        max_cols = max(len(r) for r in rows)
        if max_cols > 40:
            max_cols = 40  # Cap columns for markdown readability
        header = [str(c).replace('\n', ' ').replace('|', '\\|').strip() if c is not None else "" for c in (list(rows[0]) + [""] * max_cols)[:max_cols]]
        lines.append("| " + " | ".join(header) + " |")
        lines.append("| " + " | ".join(["---"] * max_cols) + " |")
        for r in rows[1:]:
            cells = [str(c).replace('\n', ' ').replace('|', '\\|').strip() if c is not None else "" for c in (list(r) + [""] * max_cols)[:max_cols]]
            lines.append("| " + " | ".join(cells) + " |")
        lines.append("\n")
    return "\n".join(lines)

def convert_pdf(p: Path) -> str:
    reader = PdfReader(str(p))
    lines = [f"# Document: {p.stem}\n", f"*Total Pages: {len(reader.pages)}*\n"]
    for i, page in enumerate(reader.pages):
        try:
            text = page.extract_text()
            if text and text.strip():
                lines.append(f"\n---\n*Page {i+1}*\n\n{text.strip()}\n")
        except Exception as e:
            lines.append(f"\n---\n*Page {i+1} [Extraction Note: {e}]*\n")
    return "\n".join(lines)

def convert_md(p: Path) -> str:
    return p.read_text(encoding='utf-8', errors='replace')

def main():
    files = sorted(list(SOURCE_DIR.iterdir()), key=lambda x: x.stat().st_size)
    total = len(files)
    print(f"[*] Starting conversion of {total} files from '{SOURCE_DIR}' to '{DEST_DIR}'...")
    
    success = 0
    errors = []
    t_start = time.time()
    
    for idx, f in enumerate(files, 1):
        target_name = get_target_filename(f)
        dest_path = DEST_DIR / target_name
        
        if dest_path.exists() and dest_path.stat().st_size > 0:
            success += 1
            continue

        ext = f.suffix.lower()
        size_mb = round(f.stat().st_size / (1024 * 1024), 2)
        print(f"[{idx}/{total}] Converting: {f.name} ({size_mb} MB) -> {target_name}...", end=" ", flush=True)
        
        t0 = time.time()
        try:
            if f.name.lower().endswith('.docx'):
                content = convert_docx(f)
            elif ext == '.pdf':
                content = convert_pdf(f)
            elif ext == '.xlsx':
                content = convert_xlsx(f)
            elif ext == '.md':
                content = convert_md(f)
            else:
                content = f.read_text(encoding='utf-8', errors='replace')
            
            dest_path.write_text(content, encoding='utf-8', errors='replace')
            elapsed = round(time.time() - t0, 2)
            out_kb = round(dest_path.stat().st_size / 1024, 1)
            print(f"DONE ({elapsed}s, {out_kb} KB)")
            success += 1
        except Exception as err:
            elapsed = round(time.time() - t0, 2)
            print(f"FAILED ({err})")
            errors.append((f.name, str(err)))
    
    total_time = round(time.time() - t_start, 2)
    print("\n" + "="*60)
    print(f"[*] Conversion Complete: {success}/{total} files converted in {total_time} seconds.")
    if errors:
        print(f"[!] {len(errors)} Errors encountered:")
        for name, err in errors:
            print(f"    - {name}: {err}")
    else:
        print("[+] All files converted with 100% success!")
    print("="*60)

if __name__ == "__main__":
    main()
